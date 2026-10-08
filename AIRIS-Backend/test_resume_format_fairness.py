import io
import unittest
from unittest.mock import patch

import docx
import fitz

from ai_service import extract_authoritative_resume_skills
from deterministic_scorer import evaluate_ats_machine_readability
from evidence_chunker import chunk_resume_evidence
from file_utils import extract_text
from resume_parser import parse_resume


PROFILE = [
    ("CONTACT", "Alex Rivera | alex.rivera@example.com | 09171234567 | github.com/alexrivera"),
    ("EDUCATION", "Bachelor of Science in Information Technology | Technological Institute of the Philippines | 2022 - 2026"),
    ("SKILLS", "Python, React, Firebase, SQL, Git"),
    ("CERTIFICATIONS", "AWS Cloud Practitioner | 2025"),
    ("PROJECTS", "Developed a campus scheduling application using React and Firebase in 2025. Tested with a team of five."),
    ("EXPERIENCE", "Software intern | 2025 - 2026 | Built and tested API features with Python and SQL."),
    ("ACHIEVEMENTS", "Dean's List 2024 and 2025"),
]


def make_pdf(columns=False, complex_layout=False, decorative=False, unusual=False):
    document = fitz.open()
    page = document.new_page(width=595, height=842)
    entries = list(reversed(PROFILE)) if complex_layout else PROFILE
    for index, (heading, body) in enumerate(entries):
        x = 35 + (index % 2) * 280 if columns else 40
        y = 55 + (index // 2) * 190 if columns else 60 + index * 82
        width = 260 if columns else 510
        if decorative:
            page.draw_circle((25, y - 4), 5, color=(0.2, 0.4, 0.7))
        shown_heading = {"EDUCATION": "ACADEMIC FORMATION", "SKILLS": "CAPABILITIES"}.get(heading, heading) if unusual else heading
        page.insert_text((x, y), shown_heading, fontsize=12)
        page.insert_textbox(fitz.Rect(x, y + 10, x + width, y + 105), body, fontsize=9)
    return document.tobytes()


def make_docx(mode):
    document = docx.Document()
    entries = PROFILE
    if mode == "header":
        document.sections[0].header.paragraphs[0].text = "\n".join(PROFILE[0])
        document.sections[0].footer.paragraphs[0].text = "\n".join(PROFILE[-1])
        entries = PROFILE[1:-1]
    if mode == "table":
        table = document.add_table(rows=0, cols=2)
        for heading, body in entries:
            row = table.add_row()
            row.cells[0].text = heading
            row.cells[1].text = body
    else:
        for heading, body in entries:
            document.add_paragraph(heading)
            document.add_paragraph(body)
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def score(file_bytes, filename):
    raw = extract_text(file_bytes, filename)
    parsed = parse_resume(raw)
    chunks = chunk_resume_evidence(parsed)
    skills = extract_authoritative_resume_skills(raw, parsed)
    return raw, parsed, evaluate_ats_machine_readability(chunks, parsed, skills)


class ResumeFormatFairnessTests(unittest.TestCase):
    def test_equivalent_pdf_and_docx_layouts_retain_facts_and_ats_score(self):
        variants = [
            (make_pdf(), "clean.pdf"),
            (make_pdf(columns=True), "columns.pdf"),
            (make_docx("table"), "table.docx"),
            (make_docx("header"), "header.docx"),
            (make_pdf(complex_layout=True), "complex.pdf"),
            (make_pdf(decorative=True), "decorative.pdf"),
            (make_pdf(unusual=True), "unusual.pdf"),
        ]
        results = [score(data, name) for data, name in variants]
        for raw, parsed, ats in results:
            self.assertIn("alex.rivera@example.com", raw)
            self.assertIn("Technological Institute", " ".join(raw.split()))
            self.assertIn("Firebase", raw)
            self.assertIn("ContactInformation", parsed)
            self.assertIn("Education", parsed)
            self.assertEqual(ats["overallScore"], results[0][2]["overallScore"])

    def test_image_only_pdf_uses_ocr_instead_of_binary_decoding(self):
        source = fitz.open()
        page = source.new_page()
        image = fitz.open()
        image_page = image.new_page()
        image_page.insert_text((40, 50), "EDUCATION", fontsize=12)
        png = image_page.get_pixmap().tobytes("png")
        page.insert_image(page.rect, stream=png)
        with patch("file_utils.extract_text_from_image", return_value="EDUCATION\nBachelor of Science in Information Technology") as ocr:
            extracted = extract_text(source.tobytes(), "scan.pdf")
        self.assertEqual(extracted, "EDUCATION\nBachelor of Science in Information Technology")
        ocr.assert_called_once()

    def test_corrupt_documents_do_not_become_binary_gibberish(self):
        self.assertEqual(extract_text(b"not a document" * 10, "bad.pdf"), "")
        self.assertEqual(extract_text(b"not a document" * 10, "bad.docx"), "")
        with self.assertRaisesRegex(ValueError, "Unsupported file format"):
            extract_text(b"legacy word data", "legacy.doc")

    def test_understandable_alternative_headings_are_canonicalized(self):
        raw = "ACADEMIC FORMATION\n" + PROFILE[1][1] + "\nCAPABILITIES\n" + PROFILE[2][1]
        parsed = parse_resume(raw)
        self.assertIn("Education", parsed)
        self.assertIn("Skills", parsed)


if __name__ == "__main__":
    unittest.main()
