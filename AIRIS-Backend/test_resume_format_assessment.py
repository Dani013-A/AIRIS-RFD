import unittest
from io import BytesIO
from unittest.mock import patch

import fitz
from docx import Document

from resume_format import (
    MAX_EXTRACTED_CHARACTERS, MAX_UPLOAD_BYTES, assess_resume_format,
    blocked, inspect_document_layout, validate_file_type,
)
from resume_parser import parse_resume


STANDARD = """María Dela Cruz
Email: maria.qa@example.org
EDUCATION
Polytechnic University of the Philippines
BS Information Technology, 2025
SKILLS
Python, data analysis, teamwork
PROJECTS
Capstone inventory system
Built and tested an inventory application with classmates.
"""


class ResumeFormatAssessmentTests(unittest.TestCase):
    def assess(self, text=STANDARD, filename="resume.pdf", **flags):
        return assess_resume_format(text, parse_resume(text), filename, **flags)

    def test_populated_standard_and_optional_sections_omitted(self):
        result = self.assess()
        self.assertEqual(result["outcome"], "SUPPORTED")
        self.assertTrue(result["canProceed"])
        self.assertEqual(result["specificationVersion"], "airis-resume-v1")

    def test_unfilled_template_and_placeholder_only_sections(self):
        template = """YOUR NAME
Phone: +63 9XX XXX XXXX | Email: your.email@example.com
PROFESSIONAL SUMMARY
Graduating Degree Program student preparing for entry-level opportunities as a Target Position.
EDUCATION
Institution Name | City, Philippines
Degree Program | Expected Graduation: Month Year
SKILLS
Domain skills: [List skills you can demonstrate]
PROJECTS
[Project Title] | Month Year - Month Year
[Describe your action, method, and result in plain text]
"""
        result = self.assess(template)
        self.assertEqual(result["outcome"], "BLOCKED")
        self.assertIn("PLACEHOLDER_ONLY", result["reasonCodes"])
        partial = STANDARD.replace("Python, data analysis, teamwork", "[List skills you can demonstrate]")
        result = self.assess(partial)
        self.assertEqual(result["outcome"], "WARNING")
        self.assertIn("Skills", result["missingSections"])

    def test_missing_information_is_warning_not_fabrication(self):
        for text, code in (
            (STANDARD.replace("María Dela Cruz\n", ""), "MISSING_NAME"),
            (STANDARD.replace("Email: maria.qa@example.org\n", ""), "MISSING_CONTACT"),
            (STANDARD.replace("EDUCATION\nPolytechnic University of the Philippines\nBS Information Technology, 2025\n", ""), "MISSING_EDUCATION"),
            (STANDARD.replace("SKILLS\nPython, data analysis, teamwork\n", ""), "MISSING_SKILLS"),
            (STANDARD.replace("PROJECTS\nCapstone inventory system\nBuilt and tested an inventory application with classmates.\n", ""), "MISSING_EXPERIENCE_EVIDENCE"),
        ):
            with self.subTest(code=code):
                result = self.assess(text)
                self.assertEqual(result["outcome"], "WARNING")
                self.assertIn(code, result["reasonCodes"])

    def test_aliases_reordering_extra_sections_and_role_independence(self):
        text = STANDARD.replace("EDUCATION", "ACADEMIC FORMATION").replace("SKILLS", "CAPABILITIES")
        text = text.replace("PROJECTS\nCapstone", "PROJECTS\nCapstone")
        result = self.assess(text)
        self.assertEqual(result["outcome"], "SUPPORTED")
        extra = text + "\nHOBBIES\nI practice watercolor painting.\n"
        self.assertEqual(self.assess(extra)["outcome"], "WARNING")
        self.assertIn("HOBBIES", self.assess(extra)["unrecognizedSections"])
        reordered = "SKILLS\nPython, data analysis, teamwork\n" + STANDARD.replace("SKILLS\nPython, data analysis, teamwork\n", "")
        self.assertTrue(self.assess(reordered)["canProceed"])

    def test_readable_non_pdf_and_extraction_risk(self):
        self.assertIn("NON_PDF", self.assess(filename="resume.docx")["reasonCodes"])
        self.assertIn("SCANNED_PDF", self.assess(scanned=True)["reasonCodes"])
        self.assertIn("COMPLEX_LAYOUT", self.assess(complex_layout=True)["reasonCodes"])
        self.assertEqual(self.assess(STANDARD.replace("EDUCATION", "🎨 EDUCATION"))["outcome"], "SUPPORTED")

    def test_file_validation_and_resource_bounds(self):
        self.assertIn("EXTENSION_MISMATCH", validate_file_type("x.pdf", "text/plain", b"%PDF-1.7")["reasonCodes"])
        self.assertIn("EXTENSION_MISMATCH", validate_file_type("x.pdf", "application/pdf", b"not a pdf")["reasonCodes"])
        self.assertIn("UNSUPPORTED_TYPE", validate_file_type("x.exe", "application/octet-stream", b"data")["reasonCodes"])
        self.assertIn("TEXT_TOO_LARGE", self.assess(STANDARD + "a" * MAX_EXTRACTED_CHARACTERS)["reasonCodes"])
        self.assertGreater(MAX_UPLOAD_BYTES, 0)
        self.assertIn("UNREADABLE", self.assess(" ")["reasonCodes"])

    def test_corrupt_encrypted_and_multipage_pdf(self):
        _, issue = inspect_document_layout("x.pdf", b"%PDF-corrupt")
        self.assertIn("INVALID_DOCUMENT", issue["reasonCodes"])
        document = fitz.open()
        for _ in range(3):
            page = document.new_page()
            page.insert_text((72, 72), STANDARD)
        flags, issue = inspect_document_layout("x.pdf", document.tobytes())
        self.assertIsNone(issue)
        self.assertFalse(flags["scanned"])
        encrypted = document.tobytes(encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw="owner", user_pw="user")
        _, issue = inspect_document_layout("x.pdf", encrypted)
        self.assertIn("INVALID_DOCUMENT", issue["reasonCodes"])

    def test_two_column_pdf_is_warned_and_colored_text_is_not(self):
        columns = fitz.open()
        page = columns.new_page()
        page.insert_textbox(fitz.Rect(40, 80, 270, 350),
                            "Education and university study with several substantial lines of candidate details.", fontsize=11)
        page.insert_textbox(fitz.Rect(320, 80, 560, 350),
                            "Skills and project evidence with several substantial lines of candidate details.", fontsize=11)
        flags, issue = inspect_document_layout("columns.pdf", columns.tobytes())
        self.assertIsNone(issue)
        self.assertTrue(flags["complex_layout"])
        self.assertIn("COMPLEX_LAYOUT", self.assess(**flags)["reasonCodes"])

        colored = fitz.open()
        colored.new_page().insert_text((70, 70), "María Dela Cruz", color=(0.8, 0.1, 0.5))
        flags, issue = inspect_document_layout("colored.pdf", colored.tobytes())
        self.assertIsNone(issue)
        self.assertFalse(flags["complex_layout"])

    def test_docx_table_and_header(self):
        doc = Document()
        doc.sections[0].header.paragraphs[0].text = "maria.qa@example.org"
        doc.add_table(rows=1, cols=2)
        output = BytesIO()
        doc.save(output)
        flags, issue = inspect_document_layout("x.docx", output.getvalue())
        self.assertIsNone(issue)
        self.assertTrue(flags["complex_layout"])

    def test_repeated_assessment_is_stable(self):
        self.assertEqual(self.assess(), self.assess())

    def test_upload_contract_supported_warning_and_placeholder_block(self):
        from fastapi.testclient import TestClient
        from main import app
        client = TestClient(app)
        def upload(text):
            return client.post(
                "/upload", headers={"Authorization": "Bearer qa-token"},
                files={"file": ("resume.pdf", BytesIO(b"%PDF-test"), "application/pdf")},
                data={"targetPosition": "Accountant"},
            ).json()
        analysis = {
            "resumeScore": 50, "atsScore": 50, "alignmentScore": 50,
            "overallScore": 50, "suggestions": [],
        }
        with (patch("firebase_admin.auth.verify_id_token", return_value={"uid": "qa-user", "email": "qa-user@tip.edu.ph", "email_verified": True}),
              patch("main.assert_active_student"),
              patch("main.inspect_document_layout", return_value=({}, None)),
              patch("main.extract_text", side_effect=[STANDARD, STANDARD.replace("SKILLS\nPython, data analysis, teamwork\n", ""),
                                                     "YOUR NAME\nyour.email@example.com\nInstitution Name\nDegree Program\nSKILLS\n[List skills you can demonstrate]\nPROJECTS\n[Describe your action, method, and result]"]),
              patch("main.validate_resume_validity_with_llm", return_value={"is_resume": True}),
              patch("main.evaluate_semantic_alignment", return_value={}),
              patch("main.validate_and_enhance_analysis", return_value=analysis)):
            supported = upload(STANDARD)
            warning = upload(STANDARD)
            blocked_response = upload(STANDARD)
        self.assertTrue(supported["valid"])
        self.assertEqual(supported["formatAssessment"]["outcome"], "SUPPORTED")
        self.assertEqual(warning["formatAssessment"]["outcome"], "WARNING")
        self.assertTrue(warning["valid"])
        self.assertFalse(blocked_response["valid"])
        self.assertIsNone(blocked_response["analysisId"])

    def test_populated_harvard_inspired_resume_supported_without_warning(self):
        import os
        fixture_path = os.path.join(os.path.dirname(__file__), "../scratch/investigation/Daniel_Aaron_T_Espela_AIRIS_Harvard_Inspired_Resum_261010_235100.pdf")
        if os.path.exists(fixture_path):
            doc = fitz.open(fixture_path)
            raw_text = "\n".join(page.get_text() for page in doc)
            populated_text = raw_text.replace(
                "Contact details to be added",
                "daniel.espela@tip.edu.ph | +63 917 123 4567 | linkedin.com/in/danielespela"
            )
            parsed = parse_resume(populated_text)
            flags, issue = inspect_document_layout("harvard.pdf", open(fixture_path, "rb").read())
            self.assertIsNone(issue)
            result = assess_resume_format(populated_text, parsed, "harvard.pdf", **flags)
            self.assertEqual(result["outcome"], "SUPPORTED")
            self.assertTrue(result["canProceed"])
            self.assertEqual(result["reasonCodes"], [])
            self.assertEqual(result["unrecognizedSections"], [])
            self.assertEqual(result["missingSections"], [])
            self.assertEqual(result["dimensions"]["fileValidity"]["status"], "PASS")
            self.assertEqual(result["dimensions"]["textQuality"]["status"], "PASS")
            self.assertEqual(result["dimensions"]["contentCompleteness"]["status"], "PASS")
            self.assertEqual(result["dimensions"]["formatCompatibility"]["status"], "PASS")
            self.assertEqual(result["dimensions"]["layoutConfidence"]["status"], "PASS")
            self.assertEqual(result["dimensions"]["analysisEligibility"]["status"], "SUPPORTED")
            self.assertTrue(result["dimensions"]["analysisEligibility"]["canProceed"])

    def test_harvard_inspired_resume_placeholder_contact_warns_only_contact(self):
        import os
        fixture_path = os.path.join(os.path.dirname(__file__), "../scratch/investigation/Daniel_Aaron_T_Espela_AIRIS_Harvard_Inspired_Resum_261010_235100.pdf")
        if os.path.exists(fixture_path):
            doc = fitz.open(fixture_path)
            raw_text = "\n".join(page.get_text() for page in doc)
            parsed = parse_resume(raw_text)
            flags, _ = inspect_document_layout("harvard.pdf", open(fixture_path, "rb").read())
            result = assess_resume_format(raw_text, parsed, "harvard.pdf", **flags)
            self.assertEqual(result["outcome"], "WARNING")
            self.assertTrue(result["canProceed"])
            self.assertEqual(result["reasonCodes"], ["MISSING_CONTACT"])
            self.assertEqual(result["unrecognizedSections"], [])
            self.assertEqual(result["dimensions"]["contentCompleteness"]["candidateNameDetected"], True)
            self.assertEqual(result["dimensions"]["contentCompleteness"]["contactDetected"], False)

    def test_synonyms_and_subcategories_no_unrecognized_sections(self):
        text = """JUAN DE LA CRUZ
juan@example.edu.ph | +63 912 345 6789
CAREER FOCUS
Aspiring full-stack software engineer with strong technical foundations.
ACADEMIC FORMATION
Technological Institute of the Philippines
Bachelor of Science in Information Technology, 2025
CORE STRENGTHS
Programming Languages: Python, JavaScript, Java
Frameworks & Platforms: FastAPI, React Native, Node.js
SELECTED PROJECTS
Campus Portal Application
Full stack student application with real time notifications.
"""
        result = self.assess(text)
        self.assertEqual(result["outcome"], "SUPPORTED")
        self.assertTrue(result["canProceed"])
        self.assertEqual(result["unrecognizedSections"], [])
        self.assertEqual(result["missingSections"], [])

    def test_dimensions_structure_completeness(self):
        result = self.assess()
        self.assertIn("dimensions", result)
        dims = result["dimensions"]
        self.assertIn("fileValidity", dims)
        self.assertIn("textQuality", dims)
        self.assertIn("contentCompleteness", dims)
        self.assertIn("formatCompatibility", dims)
        self.assertIn("layoutConfidence", dims)
        self.assertIn("analysisEligibility", dims)
        self.assertTrue(dims["analysisEligibility"]["canProceed"])

    def test_blocked_result_has_blocked_dimensions(self):
        result = blocked("UNSUPPORTED_TYPE")
        self.assertEqual(result["outcome"], "BLOCKED")
        self.assertFalse(result["canProceed"])
        self.assertEqual(result["dimensions"]["analysisEligibility"]["canProceed"], False)
        self.assertEqual(result["dimensions"]["fileValidity"]["status"], "FAIL")


if __name__ == "__main__":
    unittest.main()
