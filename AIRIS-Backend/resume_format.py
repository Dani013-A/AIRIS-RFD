"""Deterministic AIRIS resume-format assessment; independent of scoring and AI validity."""

import re
from pathlib import Path

SPECIFICATION_VERSION = "airis-resume-v1"
MAX_UPLOAD_BYTES = 15 * 1024 * 1024
MAX_EXTRACTED_CHARACTERS = 100_000
SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".rtf", ".txt", ".png", ".jpg", ".jpeg", ".webp", ".heic"}
_MIME_TYPES = {
    ".pdf": {"application/pdf"},
    ".docx": {"application/vnd.openxmlformats-officedocument.wordprocessingml.document", "application/zip"},
    ".rtf": {"application/rtf", "text/rtf"},
    ".txt": {"text/plain"},
    ".png": {"image/png"},
    ".jpg": {"image/jpeg"},
    ".jpeg": {"image/jpeg"},
    ".webp": {"image/webp"},
    ".heic": {"image/heic", "image/heif"},
}
_PLACEHOLDERS = (
    r"\byour name\b", r"your\.email@example\.com", r"\+63 9xx", r"\binstitution name\b",
    r"\bdegree program\b", r"\btarget position\b", r"\[list ", r"\[describe ",
    r"\[project title\]", r"\[internship, work, or relevant activity\]",
    r"\bgraduating degree program student\b",
)
_CONTACT = re.compile(
    r"[\w.+-]+@[\w.-]+\.[a-z]{2,}"
    r"|(?:\+?\d[\d\s().-]{8,}\d)"
    r"|(?:https?://)?(?:www\.)?(?:linkedin\.com/|github\.com/|gitlab\.com/|behance\.net/|dribbble\.com/)\S+"
    r"|https?://[^\s/$.?#].[^\s]*"
    r"|(?:portfolio|website|github|linkedin|gitlab):\s*\S+",
    re.I,
)
_PLACEHOLDER_CONTACT = re.compile(
    r"your\.email@example\.com|\+63 9xx|contact details to be added|details to be added",
    re.I,
)
_NAME_EXCLUSIONS = {
    "professional summary", "summary", "education", "skills", "projects", "experience",
    "contact", "contact information", "curriculum vitae", "resume",
}

_MESSAGES = {
    "UNREADABLE": "Unable to extract readable text from this document.",
    "PLACEHOLDER_ONLY": "The document still contains template examples rather than candidate information.",
    "MISSING_NAME": "A candidate name could not be recovered near the beginning of the resume.",
    "MISSING_CONTACT": "A usable email, phone number, or profile link could not be recovered.",
    "MISSING_EDUCATION": "An Education section with candidate information could not be recovered.",
    "MISSING_SKILLS": "A Skills section with candidate information could not be recovered.",
    "MISSING_EXPERIENCE_EVIDENCE": "No populated project, experience, internship, or leadership section was recovered.",
    "NON_PDF": "This readable file is accepted, but the T.I.P. standard is a text-based PDF.",
    "SCANNED_PDF": "This PDF required OCR; extraction may omit or reorder text.",
    "COMPLEX_LAYOUT": "Columns, tables, or header/footer content may be extracted out of order.",
    "UNRECOGNIZED_SECTIONS": "Some section headings were not recognized by AIRIS.",
    "EXTENSION_MISMATCH": "The file extension does not match its declared type or content.",
    "FILE_TOO_LARGE": "The file exceeds the 15 MiB processing limit.",
    "TEXT_TOO_LARGE": "The extracted text exceeds the 100,000-character processing limit.",
    "PAGE_LIMIT": "The resume exceeds the three-page limit.",
    "UNSUPPORTED_TYPE": "This file type is not supported.",
    "INVALID_DOCUMENT": "The document could not be opened or is protected.",
    "INSUFFICIENT_CONTENT": "There is not enough actual candidate information to analyze reliably.",
}


def _result(
    outcome: str,
    codes=(),
    missing=(),
    unrecognized=(),
    unrecovered=(),
    *,
    filename="resume.pdf",
    raw_text="",
    content_sections=None,
    scanned=False,
    complex_layout=False,
    page_count=1,
    extraction_method="embedded",
    is_placeholder=False,
    name_detected=True,
    contact_detected=True,
    file_issue=False,
):
    unique_codes = list(dict.fromkeys(codes))
    remediation_text = (
        "Download the T.I.P. Standard Resume Template, replace every example with your own information, "
        "and export a selectable-text PDF before retrying."
    ) if outcome != "SUPPORTED" else "No format changes required."

    dimensions = {
        "fileValidity": {
            "status": "FAIL" if (outcome == "BLOCKED" and any(c in unique_codes for c in ("UNSUPPORTED_TYPE", "EXTENSION_MISMATCH", "FILE_TOO_LARGE", "PAGE_LIMIT", "INVALID_DOCUMENT"))) else "PASS",
            "extension": Path(filename).suffix.lower(),
            "pageCount": page_count,
            "isProtected": "INVALID_DOCUMENT" in unique_codes,
            "details": "File integrity and container format check passed." if outcome != "BLOCKED" else "File integrity or container validation failed.",
        },
        "textQuality": {
            "status": "FAIL" if any(c in unique_codes for c in ("UNREADABLE", "TEXT_TOO_LARGE")) else ("WARNING" if scanned else "PASS"),
            "characterCount": len(raw_text),
            "pageCount": page_count,
            "extractionMethod": extraction_method,
            "isScanned": scanned,
            "details": "Selectable text extracted directly from document." if not scanned and len(raw_text) >= 50 else ("OCR fallback applied for scanned/image pages." if scanned else "Extraction inadequate or exceeded size limit."),
        },
        "contentCompleteness": {
            "status": "FAIL" if any(c in unique_codes for c in ("INSUFFICIENT_CONTENT", "PLACEHOLDER_ONLY")) else ("WARNING" if (missing or unrecovered) else "PASS"),
            "candidateNameDetected": name_detected,
            "contactDetected": contact_detected,
            "requiredSectionsPresent": [k for k in ("Education", "Skills", "Projects", "WorkExperience", "Internships", "OrganizationsLeadership") if content_sections and k in content_sections],
            "missingRequiredSections": list(missing),
            "unrecoveredInformation": list(unrecovered),
        },
        "formatCompatibility": {
            "status": "FAIL" if any(c in unique_codes for c in ("UNSUPPORTED_TYPE", "PLACEHOLDER_ONLY")) else ("WARNING" if (unrecognized or "NON_PDF" in unique_codes) else "PASS"),
            "isStandardAcademic": Path(filename).suffix.lower() == ".pdf" and not unrecognized and not complex_layout,
            "isPlaceholderOnly": is_placeholder,
            "recognizedSections": [k for k in (content_sections or {}).keys() if not k.startswith("Custom_") and k != "raw_text"],
            "unrecognizedSections": list(unrecognized),
            "isPdf": Path(filename).suffix.lower() == ".pdf",
        },
        "layoutConfidence": {
            "status": "WARNING" if complex_layout else "PASS",
            "complexLayout": complex_layout,
            "readingOrderReliable": not complex_layout,
            "details": "Single-column sequential layout." if not complex_layout else "Multi-column or complex structure detected; layout confidence reduced.",
        },
        "analysisEligibility": {
            "canProceed": outcome != "BLOCKED",
            "status": outcome,
            "reasonCodes": unique_codes,
            "reasons": [_MESSAGES[code] for code in unique_codes if code in _MESSAGES],
            "remediation": remediation_text,
        },
    }

    return {
        "outcome": outcome,
        "specificationVersion": SPECIFICATION_VERSION,
        "reasonCodes": unique_codes,
        "reasons": [_MESSAGES[code] for code in unique_codes if code in _MESSAGES],
        "missingSections": list(missing),
        "unrecognizedSections": list(unrecognized),
        "unrecoveredInformation": list(unrecovered),
        "remediation": remediation_text,
        "canProceed": outcome != "BLOCKED",
        "templateId": "airis_standard",
        "extractionMethod": extraction_method,
        "dimensions": dimensions,
    }


def blocked(code, filename="resume.pdf"):
    return _result("BLOCKED", [code], filename=filename, name_detected=False, contact_detected=False, file_issue=True)


def validate_file_type(filename: str, content_type: str, content: bytes):
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        return blocked("UNSUPPORTED_TYPE", filename)
    mime = (content_type or "").split(";", 1)[0].strip().lower()
    if mime and mime not in {"application/octet-stream", "binary/octet-stream"} and mime not in _MIME_TYPES[extension]:
        return blocked("EXTENSION_MISMATCH", filename)
    signatures = {
        ".pdf": content.startswith(b"%PDF-"),
        ".docx": content.startswith(b"PK\x03\x04"),
        ".rtf": content.startswith(b"{\\rtf"),
        ".png": content.startswith(b"\x89PNG\r\n\x1a\n"),
        ".jpg": content.startswith(b"\xff\xd8\xff"),
        ".jpeg": content.startswith(b"\xff\xd8\xff"),
        ".webp": content.startswith(b"RIFF") and content[8:12] == b"WEBP",
        ".heic": content[4:8] == b"ftyp" and b"hei" in content[8:16],
        ".txt": b"\x00" not in content[:4096],
    }
    if extension == ".txt" and not signatures[extension]:
        return blocked("UNREADABLE", filename)
    if not signatures[extension]:
        return blocked("EXTENSION_MISMATCH", filename)
    return None


def inspect_document_layout(filename: str, content: bytes):
    """Return extraction-risk flags, or a blocked assessment for invalid containers."""
    extension = Path(filename).suffix.lower()
    if extension == ".pdf":
        try:
            import fitz
            document = fitz.open(stream=content, filetype="pdf")
            if document.is_encrypted or not document.page_count:
                return {}, blocked("INVALID_DOCUMENT", filename)
            page_count = document.page_count
            scanned = False
            columns = False
            for page in document:
                blocks = [b for b in page.get_text("blocks") if len(b) >= 5 and str(b[4]).strip()]
                if len(page.get_text().strip()) < 30 and page.get_images():
                    scanned = True
                # Side-by-side substantial text blocks with overlapping vertical spans.
                for index, left in enumerate(blocks):
                    for right in blocks[index + 1:]:
                        if (len(str(left[4]).strip()) > 35 and len(str(right[4]).strip()) > 35
                                and left[2] < right[0] and min(left[3], right[3]) - max(left[1], right[1]) > 12):
                            columns = True
            return {
                "scanned": scanned,
                "complex_layout": columns,
                "page_count": page_count,
                "extraction_method": "ocr" if scanned else "embedded",
            }, None
        except Exception:
            return {}, blocked("INVALID_DOCUMENT", filename)
    if extension == ".docx":
        try:
            from io import BytesIO
            from zipfile import ZipFile
            from docx import Document
            with ZipFile(BytesIO(content)) as archive:
                if len(archive.infolist()) > 1000 or sum(info.file_size for info in archive.infolist()) > 50 * 1024 * 1024:
                    return {}, blocked("FILE_TOO_LARGE", filename)
            document = Document(BytesIO(content))
            has_header_footer = any(
                any(p.text.strip() for p in part.paragraphs)
                for section in document.sections for part in (section.header, section.footer)
            )
            return {
                "complex_layout": bool(document.tables) or has_header_footer,
                "page_count": 1,
                "extraction_method": "embedded",
            }, None
        except Exception:
            return {}, blocked("INVALID_DOCUMENT", filename)
    return {
        "scanned": False,
        "complex_layout": False,
        "page_count": 1,
        "extraction_method": "ocr" if extension in {".png", ".jpg", ".jpeg", ".webp", ".heic"} else "embedded",
    }, None


def assess_resume_format(
    raw_text: str,
    sections: dict,
    filename: str,
    *,
    scanned=False,
    complex_layout=False,
    page_count=1,
    extraction_method="embedded",
    **kwargs
):
    if not raw_text or len(raw_text.strip()) < 30:
        return blocked("UNREADABLE", filename)
    if len(raw_text) > MAX_EXTRACTED_CHARACTERS:
        return blocked("TEXT_TOO_LARGE", filename)

    def meaningful(value):
        clean = "\n".join(
            line for line in (value or "").splitlines()
            if not any(re.search(phrase, line, re.I) for phrase in _PLACEHOLDERS)
        )
        clean = re.sub(r"\[[^\]]{1,120}\]", "", clean)
        clean = re.sub(r"\b(?:Month Year|City, Philippines|Expected Graduation)\b", "", clean, flags=re.I)
        clean = re.sub(r"[^\w\s]", " ", clean, flags=re.UNICODE)
        return len(re.findall(r"[^\W\d_]{2,}", clean, flags=re.UNICODE)) >= 2

    content_sections = {key: value for key, value in sections.items() if key != "raw_text" and meaningful(value)}
    from resume_parser import NORMALIZED_ALIASES, SKILL_SUBCATEGORY_PREFIXES, _normalize_heading

    first_lines = []
    header_lines = set()
    for line in raw_text.splitlines():
        candidate = line.strip()
        if not candidate:
            continue
        normalized = _normalize_heading(candidate)
        if normalized in NORMALIZED_ALIASES and NORMALIZED_ALIASES[normalized] != "ContactInformation":
            break
        header_lines.add(candidate.lower())
        if len(first_lines) < 8:
            first_lines.append(candidate.split("|", 1)[0].strip())

    name = any(
        2 <= len(line.split()) <= 6 and len(line) <= 80
        and line.lower() not in _NAME_EXCLUSIONS
        and not _CONTACT.search(line)
        and not any(re.search(phrase, line, re.I) for phrase in _PLACEHOLDERS)
        and not line.startswith(("#", "[", "•", "-", "*"))
        and not re.search(r"\b(?:city|philippines|phone|email|location|contact|address)\b", line, re.I)
        for line in first_lines
    )

    contact_matches = _CONTACT.findall(raw_text)
    genuine_contact = [
        m for m in contact_matches
        if not _PLACEHOLDER_CONTACT.search(m)
    ]
    contact = bool(genuine_contact)

    evidence_keys = ("Projects", "WorkExperience", "Internships", "OrganizationsLeadership")
    core_present = [key for key in ("Education", "Skills") if key in content_sections]
    evidence_present = any(key in content_sections for key in evidence_keys)
    placeholder_count = sum(bool(re.search(phrase, raw_text, re.I)) for phrase in _PLACEHOLDERS)

    if placeholder_count >= 5 or (placeholder_count >= 3 and len(core_present) < 2 and not evidence_present):
        return blocked("PLACEHOLDER_ONLY", filename)
    if not core_present and not evidence_present:
        return blocked("INSUFFICIENT_CONTENT", filename)

    missing = [key for key in ("Education", "Skills") if key not in content_sections]
    if not evidence_present:
        missing.append("Projects or Experience")
    unrecovered = []
    codes = []
    if not name:
        codes.append("MISSING_NAME")
        unrecovered.append("candidate name")
    if not contact:
        codes.append("MISSING_CONTACT")
        unrecovered.append("contact details")
    for section in missing:
        codes.append("MISSING_EXPERIENCE_EVIDENCE" if section == "Projects or Experience" else f"MISSING_{section.upper()}")
        unrecovered.append(section.lower())

    unrecognized = [key.removeprefix("Custom_") for key in sections if key.startswith("Custom_")]
    for line in raw_text.splitlines():
        heading = line.strip()
        if not heading or heading.lower() in header_lines:
            continue
        if not re.match(r"^[A-Za-z\s&/-]+$", heading):
            continue
        words = heading.split()
        if not (1 <= len(words) <= 4 and 2 <= len(heading) <= 40):
            continue
        if not heading.isupper():
            continue
        norm = _normalize_heading(heading)
        if norm in NORMALIZED_ALIASES:
            continue
        if heading.lower() in SKILL_SUBCATEGORY_PREFIXES or any(heading.lower().startswith(p) for p in SKILL_SUBCATEGORY_PREFIXES):
            continue
        if re.search(r"\b(?:bachelor|master|doctor|bs|ba|ms|phd|associate|diploma|senior high|high school|grade)\b", heading, re.I):
            continue
        if re.search(r"\b(?:university|college|institute|school|academy|campus|deped|ched)\b", heading, re.I):
            continue
        if re.search(r"\b(?:student|intern|developer|engineer|specialist|officer|assistant|lead|manager|analyst|programmer)\b", heading, re.I):
            continue
        if re.search(r"\b(?:philippines|manila|quezon|makati|pasig|city|street|avenue)\b", heading, re.I):
            continue
        if heading not in unrecognized:
            unrecognized.append(heading)

    if unrecognized:
        codes.append("UNRECOGNIZED_SECTIONS")
    if Path(filename).suffix.lower() != ".pdf":
        codes.append("NON_PDF")
    if scanned:
        codes.append("SCANNED_PDF")
    if complex_layout:
        codes.append("COMPLEX_LAYOUT")

    return _result(
        "WARNING" if codes else "SUPPORTED",
        codes,
        missing,
        unrecognized,
        unrecovered,
        filename=filename,
        raw_text=raw_text,
        content_sections=content_sections,
        scanned=scanned,
        complex_layout=complex_layout,
        page_count=page_count,
        extraction_method=extraction_method,
        is_placeholder=False,
        name_detected=name,
        contact_detected=contact,
    )
